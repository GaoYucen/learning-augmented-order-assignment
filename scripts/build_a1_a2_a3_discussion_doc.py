from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "outputs" / "A1_A2_A3_deliverables_20260820"
DOCX_PATH = OUT_DIR / "A1_A2_A3_experiment_discussion.docx"
MD_PATH = OUT_DIR / "A1_A2_A3_experiment_discussion.md"


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def fmt(x: str | float, digits: int = 3) -> str:
    try:
        return f"{float(x):.{digits}f}"
    except Exception:
        return str(x)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header, bold=True)
        set_cell_shading(table.rows[0].cells[idx], "F2F4F7")
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(10.5)


def add_para(doc: Document, text: str, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        r1.font.name = "Calibri"
        r1._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        r1.font.size = Pt(10.5)
        rest = text[len(bold_prefix) :]
        r2 = p.add_run(rest)
        r2.font.name = "Calibri"
        r2._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        r2.font.size = Pt(10.5)
    else:
        run = p.add_run(text)
        run.font.name = "Calibri"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(10.5)


def set_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.1
    normal.paragraph_format.space_after = Pt(6)
    for name, size, color in [
        ("Heading 1", 16, "2E74B5"),
        ("Heading 2", 13, "2E74B5"),
        ("Heading 3", 12, "1F4D78"),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)


def build_docx() -> None:
    summary_path = OUT_DIR / "A1_A2_synthetic_pipeline" / "synthetic_sanity_summary.csv"
    frontier_path = OUT_DIR / "A3_consistency_robustness" / "consistency_robustness_frontier.csv"
    a2_png = OUT_DIR / "A1_A2_synthetic_pipeline" / "alg_over_opt_vs_prediction_error.png"
    a3_png = OUT_DIR / "A3_consistency_robustness" / "consistency_robustness_frontier.png"

    summary = read_csv(summary_path)
    frontier = read_csv(frontier_path)

    accurate = [r for r in summary if r["prediction_scale"] == "1.0" and r["prediction_corruption"] == "scale"]
    under_05 = [r for r in summary if r["prediction_scale"] == "0.5"]
    over_15 = [r for r in summary if r["prediction_scale"] == "1.5"]

    doc = Document()
    set_styles(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("A1/A2/A3 阶段性实验结果讨论稿")
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(22)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string("0B2545")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(10)
    r = subtitle.add_run("学生 A：呼杨柯 | 分支：hyk-core-algorithms | 日期：2026-08-27")
    r.font.name = "Calibri"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string("555555")

    doc.add_heading("1. 一句话结论", level=1)
    add_para(
        doc,
        "目前 A1/A2/A3 已经形成了可运行的学习增强订单分配实验框架，并能在合成预测误差下观察到 theta 带来的 consistency-robustness trade-off。"
        "这里 theta 已统一为鲁棒 IPD 分支资源比例：theta=1 表示完全不信预测，theta=0 表示完全依赖预测 advice。"
        "但是这些结果仍基于 synthetic prediction / artificial corruption，尚未接入师姐的真实预测器，因此只能作为框架验证和讨论依据，不能作为最终论文结论。",
    )

    doc.add_heading("2. 目前完成了什么", level=1)
    add_bullet(doc, "A1：统一了 request type、资源约束、feasible bus、offline OPT、prediction error、advice error 以及算法调用接口。")
    add_bullet(doc, "A2：建立 synthetic sanity-check，用可控预测误差检查 Prediction-only、IPD、RP-LAIPD 等方法是否能跑通。")
    add_bullet(doc, "A3：建立 consistency-robustness frontier 实验，扫描 theta，观察预测准确时的收益和预测错误时的退化。")
    add_bullet(doc, "RP-LAIPD 已统一为老师口径和经典 learning-augmented 文献习惯：theta 份资源给 robust IPD branch，1-theta 份资源给 advice branch。")

    doc.add_heading("3. A2：Synthetic Sanity Check 结果", level=1)
    add_para(doc, "核心指标是 alg_over_opt，表示算法收益 / offline OPT 上界。数值越接近 1，说明越接近离线最优。")

    a2_rows = []
    for label, group in [("预测准确 scale=1.0", accurate), ("低估 scale=0.5", under_05), ("高估 scale=1.5", over_15)]:
        for method in ["IPD", "Prediction-only", "RP-LAIPD"]:
            candidates = [r for r in group if r["method"] == method]
            if method == "RP-LAIPD":
                for theta in ["0.2", "0.4", "0.6", "0.8"]:
                    r = next((x for x in candidates if x["theta"] == theta), None)
                    if r:
                        a2_rows.append([label, f"RP-LAIPD theta={theta} (robust)", fmt(r["alg_over_opt_mean"]), fmt(r["served_orders_mean"], 1), fmt(r["high_value_rejected_mean"], 1)])
            elif candidates:
                r = candidates[0]
                a2_rows.append([label, method, fmt(r["alg_over_opt_mean"]), fmt(r["served_orders_mean"], 1), fmt(r["high_value_rejected_mean"], 1)])

    add_table(doc, ["场景", "方法", "alg/OPT", "服务订单数", "高价值拒单数"], a2_rows)
    add_para(
        doc,
        "观察：预测准确时，Prediction-only 通常更有优势；theta 越小，RP-LAIPD 越接近预测 advice；theta 越大，RP-LAIPD 越接近 IPD。"
        "这符合学习增强算法的直觉：更少保留鲁棒资源可以换来更好的 consistency，但会牺牲预测失准时的 robustness。",
    )
    if a2_png.exists():
        doc.add_picture(str(a2_png), width=Inches(6.2))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("4. A3：Consistency-Robustness Frontier 结果", level=1)
    add_para(doc, "A3 固定算法结构，扫描 theta。theta=1 近似退回 IPD；theta=0 近似 Prediction-only；theta 越大，算法越保守、越不依赖 prediction advice。")
    frontier_rows = [
        [
            fmt(r["theta"], 1),
            fmt(r["consistency"]),
            fmt(r["robustness"]),
            r["worst_scenario"],
            fmt(r["consistency_ci95"]),
        ]
        for r in frontier
    ]
    add_table(doc, ["theta", "consistency", "robustness", "最坏场景", "95% CI"], frontier_rows)
    add_para(
        doc,
        "观察：theta 越小，算法越相信预测，预测准确时 consistency 更高；theta 越大，算法越接近 IPD，预测错误时 robustness 更高。"
        "这说明当前实验已经能画出清晰的 trade-off 曲线。问题也很明显：hard resource partition 较保守，在强扰动下会把资源锁在错误 advice 上，导致低 theta 的鲁棒性下降。",
    )
    if a3_png.exists():
        doc.add_picture(str(a3_png), width=Inches(6.2))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("5. 当前结果应该怎么理解", level=1)
    add_para(doc, "可以说：", bold_prefix="可以说：")
    add_bullet(doc, "A1/A2/A3 的实验框架已经初步搭好，算法接口、预测接口、误差指标和输出文件都能串起来。")
    add_bullet(doc, "RP-LAIPD 的 theta 确实控制了“保留 IPD 鲁棒性”和“相信预测 advice”之间的权衡。")
    add_bullet(doc, "当前实验可以作为周会讨论依据，用来确认实验设计、指标定义和后续接入真实预测器的接口。")
    add_para(doc, "不能说：", bold_prefix="不能说：")
    add_bullet(doc, "不能说最终算法已经优于所有 baseline，因为当前还没接入真实预测器，也没完成 A4 的完整 baseline comparison。")
    add_bullet(doc, "不能把 synthetic prediction 的结果当作真实成都订单数据上的最终效果。")
    add_bullet(doc, "不能直接声称算法创新完全来自我们；更准确的说法是：我们在实现老师期刊草稿中的学习增强 IPD 方向，并搭建验证框架。")

    doc.add_heading("6. 主要风险和需要讨论的问题", level=1)
    add_bullet(doc, "当前 synthetic setting 比较 harsh，RP-LAIPD 在部分 theta 下弱于 IPD，需要判断是实验负载设置问题，还是 hard partition 设计本身过保守。")
    add_bullet(doc, "如果真实预测器误差较小，较小的 theta 可能更有优势；如果误差较大，需要增大 theta 或考虑 adaptive theta / resource release。")
    add_bullet(doc, "A4 需要补全 Random、Greedy、FPD、IPD、Prediction-only、RP-LAIPD 的整体比较，并在真实预测器输出上重跑。")
    add_bullet(doc, "代码本地分支目前 ahead 1，最新提交还需要网络正常时 push 到 GitHub。")

    doc.add_heading("7. 建议下一步", level=1)
    add_bullet(doc, "短期：先不要合并师姐分支，保持 predictor adapter 接口，等周五讨论确认真实预测输出格式。")
    add_bullet(doc, "接入：用师姐预测器输出的 slot_id/type_id/predicted_count 替换 synthetic prediction，重跑 A2/A3。")
    add_bullet(doc, "扩展：完成 A4 overall baseline comparison，统一表格和图。")
    add_bullet(doc, "改进：若 RP-LAIPD 仍不稳定，讨论 adaptive resource release 或根据 prediction confidence 自动调 theta。")

    doc.add_heading("8. 文件位置", level=1)
    add_para(
        doc,
        f"交付物主文件夹：{OUT_DIR}。其中 A1_A2_synthetic_pipeline 保存 A2 的 CSV 和图，"
        "A3_consistency_robustness 保存 A3 的 CSV 和 frontier 图；详细路径可看同目录 Markdown 版。",
    )

    doc.save(DOCX_PATH)


def build_markdown() -> None:
    summary_path = OUT_DIR / "A1_A2_synthetic_pipeline" / "synthetic_sanity_summary.csv"
    frontier_path = OUT_DIR / "A3_consistency_robustness" / "consistency_robustness_frontier.csv"
    summary = read_csv(summary_path)
    frontier = read_csv(frontier_path)
    frontier_md_rows = "\n".join(
        f"| {fmt(row['theta'], 1)} | {fmt(row['consistency'])} | {fmt(row['robustness'])} | {row['worst_scenario']} |"
        for row in frontier
    )
    text = f"""# A1/A2/A3 阶段性实验结果讨论稿

学生 A：呼杨柯  
分支：hyk-core-algorithms  
日期：2026-08-27

## 一句话结论

目前 A1/A2/A3 已经形成了可运行的学习增强订单分配实验框架，并能在合成预测误差下观察到 theta 带来的 consistency-robustness trade-off。
这里 theta 已按导师新草稿统一为鲁棒 IPD 分支资源比例：theta=1 表示完全不信预测并退化为 robust IPD，theta=0 表示完全依赖预测 advice。
但是这些结果仍基于 synthetic prediction / artificial corruption，尚未接入师姐的真实预测器，因此只能作为框架验证和讨论依据，不能作为最终论文结论。

## 目前完成了什么

- A1：统一了 request type、资源约束、feasible bus、offline OPT、prediction error、advice error 以及算法调用接口。
- A2：建立 synthetic sanity-check，用可控预测误差检查 Prediction-only、IPD、RP-LAIPD 等方法是否能跑通。
- A3：建立 consistency-robustness frontier 实验，扫描 theta，观察预测准确时的收益和预测错误时的退化。
- RP-LAIPD 已统一为导师新草稿的符号：theta 份资源给 robust IPD branch，1-theta 份资源给 advice branch。

## 关键实验结果

### A2 Synthetic Sanity Check

- 预测准确时：Prediction-only 通常更有优势；较小 theta 的 RP-LAIPD 更接近预测 advice。
- 预测严重错误时：较大 theta 的 RP-LAIPD 更接近 IPD，鲁棒性更强。
- 解释：theta 越小越相信预测，theta 越大越保守、越接近 IPD。

图：{OUT_DIR / "A1_A2_synthetic_pipeline" / "alg_over_opt_vs_prediction_error.png"}

### A3 Consistency-Robustness Frontier

| theta | consistency | robustness | 最坏场景 |
|---:|---:|---:|---|
{frontier_md_rows}

解释：theta 越小，预测准确时收益越高；theta 越大，预测严重错误时鲁棒性越高。

图：{OUT_DIR / "A3_consistency_robustness" / "consistency_robustness_frontier.png"}

## 当前结果应该怎么说

可以说：

- A1/A2/A3 的实验框架已经初步搭好，算法接口、预测接口、误差指标和输出文件都能串起来。
- RP-LAIPD 的 theta 确实控制了“保留 IPD 鲁棒性”和“相信预测 advice”之间的权衡。
- 当前实验可以作为周会讨论依据，用来确认实验设计、指标定义和后续接入真实预测器的接口。

不能说：

- 不能说最终算法已经优于所有 baseline，因为当前还没接入真实预测器，也没完成 A4 的完整 baseline comparison。
- 不能把 synthetic prediction 的结果当作真实成都订单数据上的最终效果。
- 不能直接声称算法创新完全来自我们；更准确的说法是：我们在实现老师期刊草稿中的学习增强 IPD 方向，并搭建验证框架。

## 主要风险和需要讨论的问题

- 当前 synthetic setting 比较 harsh，RP-LAIPD 在部分 theta 下弱于 IPD，需要判断是实验负载设置问题，还是 hard partition 设计本身过保守。
- 如果真实预测器误差较小，较小的 theta 可能更有优势；如果误差较大，需要增大 theta 或考虑 adaptive theta / resource release。
- A4 需要补全 Random、Greedy、FPD、IPD、Prediction-only、RP-LAIPD 的整体比较，并在真实预测器输出上重跑。
- 代码本地分支目前 ahead 1，最新提交还需要网络正常时 push 到 GitHub。

## 建议下一步

1. 先不要合并师姐分支，保持 predictor adapter 接口，等周五讨论确认真实预测输出格式。
2. 用师姐预测器输出的 slot_id/type_id/predicted_count 替换 synthetic prediction，重跑 A2/A3。
3. 完成 A4 overall baseline comparison，统一表格和图。
4. 若 RP-LAIPD 仍不稳定，讨论 adaptive resource release 或根据 prediction confidence 自动调 theta。

## 文件位置

- 交付物文件夹：{OUT_DIR}
- A2 summary CSV：{OUT_DIR / "A1_A2_synthetic_pipeline" / "synthetic_sanity_summary.csv"}
- A3 frontier CSV：{OUT_DIR / "A3_consistency_robustness" / "consistency_robustness_frontier.csv"}
"""
    MD_PATH.write_text(text, encoding="utf-8")


if __name__ == "__main__":
    build_docx()
    build_markdown()
    print(DOCX_PATH)
    print(MD_PATH)
