from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
import pandas as pd

root = Path(r"C:\Users\hyk\Documents\learning-augmented-order-assignment")
out_dir = root / "outputs" / "report"
out_dir.mkdir(parents=True, exist_ok=True)
report_path = out_dir / "A2_A4_and_chengdu_figures_report.docx"

# Inputs
fig_a2 = root / "outputs" / "a2_sanity_check" / "alg_over_opt_vs_prediction_error.png"
fig_a4 = root / "outputs" / "chengdu" / "a4_baseline_comparison" / "a4_baseline_comparison.png"
fig_1 = root / "outputs" / "chengdu" / "final_figures" / "figure_1_model_accuracy_dispatch.png"
fig_2 = root / "outputs" / "chengdu" / "final_figures" / "figure_2_station_prediction_and_theta.png"
a4_summary = pd.read_csv(root / "outputs" / "chengdu" / "a4_baseline_comparison" / "a4_baseline_summary.csv")
a2_summary = pd.read_csv(root / "outputs" / "a2_sanity_check" / "synthetic_sanity_summary.csv")


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top', 'left', 'bottom', 'right'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f'w:{edge}'
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key in ["val", "sz", "space", "color"]:
                if key in edge_data:
                    element.set(qn(f'w:{key}'), str(edge_data[key]))


def style_table(table):
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
            set_cell_border(cell,
                top={"val": "single", "sz": 6, "color": "D9D9D9"},
                bottom={"val": "single", "sz": 6, "color": "D9D9D9"},
                left={"val": "single", "sz": 6, "color": "D9D9D9"},
                right={"val": "single", "sz": 6, "color": "D9D9D9"},
            )


def format_run(run, *, bold=False, size=11, name='Microsoft YaHei'):
    run.bold = bold
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:ascii'), name)
    run._element.rPr.rFonts.set(qn('w:hAnsi'), name)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_paragraph(doc, text, style=None, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        format_run(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        format_run(r2)
    else:
        r = p.add_run(text)
        format_run(r)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(text)
    format_run(r)
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    format_run(r, bold=True, size=14 if level == 1 else 12)
    return p


def add_table(doc, df, cols, title=None, max_rows=None):
    if title:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(title)
        format_run(r, bold=True, size=11)
    if max_rows:
        df = df.head(max_rows)
    table = doc.add_table(rows=1, cols=len(cols))
    style_table(table)
    hdr = table.rows[0].cells
    for i, (col, width) in enumerate(cols):
        hdr[i].text = col
        for p in hdr[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                format_run(run, bold=True, size=10)
        set_cell_shading(hdr[i], 'D9E2F3')
        hdr[i].width = width
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for i, (col, width) in enumerate(cols):
            val = row[col]
            if isinstance(val, float):
                text = f"{val:.3f}"
            else:
                text = str(val)
            cells[i].text = text
            cells[i].width = width
            for p in cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i != 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    format_run(run, size=10)
    return table


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.85)
section.bottom_margin = Inches(0.85)
section.left_margin = Inches(0.9)
section.right_margin = Inches(0.9)

# Base font
styles = doc.styles
styles['Normal'].font.name = 'Microsoft YaHei'
styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Microsoft YaHei')
styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Microsoft YaHei')
styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
styles['Normal'].font.size = Pt(10.5)

# Title
p = doc.add_paragraph(style='Title')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(8)
r = p.add_run('A2 A4 实验图表总结报告')
format_run(r, bold=False, size=18)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(12)
r = p.add_run('学习增强订单分配项目实验汇总')
format_run(r, size=11)

add_paragraph(doc, '本文总结当前仓库中已经生成的主要实验图和配套结果，覆盖 A2 的 synthetic sanity-check、A4 的基线对比，以及成都真实数据上的预测与 theta 分析。结论先说：预测器已经接入，A2/A4 的骨架已经补齐，但成都数据上 RP-LAIPD 还没有压过所有传统基线，说明这组实验更适合用来展示趋势和机制，而不是直接宣称最优。')

add_heading(doc, '1 实验范围', 1)
add_bullet(doc, 'A2：可控 synthetic bottleneck 场景下的 ALG/OPT 与预测误差关系。')
add_bullet(doc, 'A4：成都真实数据上的完整 baseline comparison，包含 Random、Greedy、IPD、Prediction-only、Static LP / bid-price、RP-LAIPD 和 Offline OPT。')
add_bullet(doc, '补充图：成都预测准确率梯度图，以及站点级预测与 theta 灵敏度图。')

add_heading(doc, '2 A2 synthetic sanity-check', 1)
add_paragraph(doc, '图 A2 展示了在人工可控的预测噪声下，不同方法的 ALG/OPT 变化。这个图的目的不是追求真实业务最优，而是检查方法对预测误差的响应是否符合预期。')

if fig_a2.exists():
    doc.add_picture(str(fig_a2), width=Inches(6.8))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

best_a2 = (
    a2_summary.loc[a2_summary['method'].isin(['IPD', 'Prediction-only', 'RP-LAIPD'])]
    .sort_values(['prediction_scale','method','theta'])
)
best_a2 = best_a2.loc[
    best_a2['prediction_scale'].isin([0.5, 1.0, 1.5])
    & (
        (best_a2['method'].eq('IPD'))
        | (best_a2['method'].eq('Prediction-only'))
        | ((best_a2['method'].eq('RP-LAIPD')) & best_a2['theta'].isin([0.2, 0.4, 0.6, 0.8]))
    )
].copy()
best_a2 = best_a2.rename(
    columns={
        'prediction_scale': 'scale',
        'corruption_strength': 'noise',
        'alg_over_opt_mean': 'ALG/OPT',
    }
)
add_table(
    doc,
    best_a2[['method','theta','scale','noise','ALG/OPT']].head(10),
    [
        ('method', Inches(1.7)),
        ('theta', Inches(0.7)),
        ('scale', Inches(0.95)),
        ('noise', Inches(0.85)),
        ('ALG/OPT', Inches(1.05)),
    ],
    title='表 1 A2 代表性结果',
)
add_paragraph(doc, 'A2 的主要观察是：预测误差变大后，Prediction-only 的下降更明显；IPD 更稳，但上限通常不如有较好预测的 RP-LAIPD；RP-LAIPD 在中等 theta 区间更平衡，说明 theta 的作用是把鲁棒性和预测收益分开调节。')

add_heading(doc, '3 A4 baseline comparison', 1)
add_paragraph(doc, 'A4 直接回答“当前实现和传统基线相比处于什么位置”。这里把各方法的平均表现、离线最优、运行时间和约束违例都放在一起看。')

if fig_a4.exists():
    doc.add_picture(str(fig_a4), width=Inches(6.8))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_paragraph(doc, '图 A4：左图比较不同算法的 ALG/OPT，右图比较 accepted value。误差线为 95% confidence interval，用来观察结果波动。')

baseline = a4_summary[['method','theta','accepted_value_mean','alg_over_opt_mean','served_orders_mean','slot_runtime_ms_mean']].copy()
baseline = baseline.loc[baseline['method'].isin(['Random', 'Greedy', 'IPD', 'Prediction-only', 'Static LP / bid-price', 'RP-LAIPD', 'Offline OPT'])]
baseline = baseline.rename(
    columns={
        'accepted_value_mean': 'value',
        'alg_over_opt_mean': 'ALG/OPT',
        'served_orders_mean': 'served',
        'slot_runtime_ms_mean': 'ms',
    }
)
add_table(
    doc,
    baseline[['method','theta','value','ALG/OPT','served','ms']],
    [
        ('method', Inches(1.9)),
        ('theta', Inches(0.7)),
        ('value', Inches(1.1)),
        ('ALG/OPT', Inches(0.95)),
        ('served', Inches(0.95)),
        ('ms', Inches(0.8)),
    ],
    title='表 2 A4 基线对比摘要',
)
add_paragraph(doc, '这组结果里，Greedy 和 Random 都比静态 LP / bid-price 更强；IPD 和 RP-LAIPD 处于中间位置；Offline OPT 作为上界最好。也就是说，这一版更重要的价值是把实验接口和对比框架统一起来，而不是证明成都数据上某个方法已经显著胜出。')

add_heading(doc, '4 成都预测与 theta 图', 1)
add_paragraph(doc, '这两张图是给老师讨论时最容易直接看的部分：第一张看预测器准确率梯度，第二张看站点级预测与 theta 变化下的调度敏感性。')

if fig_1.exists():
    doc.add_picture(str(fig_1), width=Inches(6.8))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_paragraph(doc, '图 1：不同预测模型版本的 Test WAPE 对比。HGB checkpoint 从低迭代到高迭代，误差逐步下降，说明可以自然构造预测准确率梯度。')

if fig_2.exists():
    doc.add_picture(str(fig_2), width=Inches(6.8))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_paragraph(doc, '图 2：左侧是代表性站点的真实需求与多模型预测，右侧是不同 theta 下 ALG/OPT 随预测误差的变化。theta=0 对应完全相信预测，theta=1 对应完全稳健。')

add_heading(doc, '5 可直接用于汇报的结论', 1)
add_bullet(doc, '预测器已经独立接入，并且能形成稳定的准确率梯度。')
add_bullet(doc, 'A2 的 synthetic 图验证了方法对预测误差的响应方向是合理的。')
add_bullet(doc, 'A4 已补上完整 baseline 框架，但成都真实数据上并没有出现“所有情况下 RP-LAIPD 都碾压基线”的结果。')
add_bullet(doc, '当前更合适的表述是：通过调 theta，RP-LAIPD 在不同预测质量下可以保持较稳的折中表现。')

add_paragraph(doc, '附：所有原始结果已同步保存在 outputs/a2_sanity_check、outputs/chengdu/a4_baseline_comparison 和 outputs/chengdu/final_figures 目录下，便于后续重新出图或改图。')

doc.save(str(report_path))
print(report_path)
